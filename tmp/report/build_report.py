from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"D:/Document Management")
OUT = ROOT / "output" / "documents" / "รายงานสหกิจศึกษา_Document_Management_System_ธนภัทร.docx"
PAGE_MAP_FILE = ROOT / "tmp" / "report" / "page_map.json"
PAGE_MAP = json.loads(PAGE_MAP_FILE.read_text(encoding="utf-8")) if PAGE_MAP_FILE.exists() else {}
FONT = "TH SarabunPSK"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.1), Cm(2.0)
sec.left_margin, sec.right_margin = Cm(2.55), Cm(2.15)
sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)


def set_font(style, size=16, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = None
    rpr = style.element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{name}"), FONT)


normal = doc.styles["Normal"]
set_font(normal)
normal.paragraph_format.line_spacing = 1.05
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.widow_control = True

for style_name, size in (("Title", 20), ("Heading 1", 18), ("Heading 2", 16), ("Heading 3", 16)):
    st = doc.styles[style_name]
    set_font(st, size, True)
    st.paragraph_format.space_before = Pt(4)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.keep_with_next = True


def paragraph(text="", align=None, indent=False, bold=False, style=None, before=0, after=0, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = FONT
        r.font.size = Pt(16)
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    return p


def spacer(pt=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pt)
    p.add_run(" ")
    return p


def page_break():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def top_letter(letter, title="", new_page=True):
    if new_page:
        page_break()
    paragraph(letter, align=WD_ALIGN_PARAGRAPH.RIGHT, after=6)
    if title:
        paragraph(title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=10)


def h(text, level=2):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def body(text):
    return paragraph(text, indent=True, after=2)


def bullet(text):
    p = paragraph("- " + text, after=1)
    p.paragraph_format.left_indent = Cm(1.05)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    return p


def cell_text(cell, value, bold=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(value)
    r.bold = bold
    r.font.name = FONT
    r.font.size = Pt(14)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def borders(table):
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            b = tcPr.first_child_found_in("w:tcBorders")
            if b is None:
                b = OxmlElement("w:tcBorders")
                tcPr.append(b)
            for edge in ("top", "left", "bottom", "right"):
                tag = qn("w:" + edge)
                node = b.find(tag)
                if node is None:
                    node = OxmlElement("w:" + edge)
                    b.append(node)
                node.set(qn("w:val"), "single")
                node.set(qn("w:sz"), "4")
                node.set(qn("w:color"), "777777")
            margins = tcPr.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tcPr.append(margins)
            for side in ("top", "bottom", "left", "right"):
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), "85")
                el.set(qn("w:type"), "dxa")
                margins.append(el)


TABLES = []
FIGURES = []
APP_FIGURES = []


def table(title, headers, rows, widths=None):
    number = len(TABLES) + 1
    caption = f"ตารางที่ {number} {title}"
    TABLES.append(caption)
    paragraph(caption, before=7, after=3, keep=True)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
    for i, head in enumerate(headers):
        cell_text(t.rows[0].cells[i], head, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for values in rows:
        cells = t.add_row().cells
        for i, val in enumerate(values):
            cell_text(cells[i], str(val))
    borders(t)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_together = True
    return t


def figure(title, instruction, appendix=False):
    arr = APP_FIGURES if appendix else FIGURES
    label = f"ภาพภาคผนวกที่ ก{len(arr)+1}" if appendix else f"ภาพที่ {len(arr)+1}"
    caption = f"{label} {title}"
    arr.append(caption)
    diagram_files = {
        1: "01_as_is_to_be.png",
        2: "02_use_case_dms.png",
        4: "01_flow_create_submit.png",
        5: "02_flow_approval.png",
        6: "06_sequence_create.png",
        7: "07_sequence_approval_signature.png",
        8: "08_er_core_prisma.png",
    }
    diagram = None if appendix else diagram_files.get(len(arr))
    if diagram:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ROOT / "output" / "diagrams" / diagram), width=Cm(15.4))
    else:
        p = paragraph(f"(นำภาพ{instruction}มาติด)", align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=6, keep=True)
        p.paragraph_format.space_before = Pt(18)
    paragraph(caption, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)


def toc_line(title, key, indent=0):
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    t.columns[0].width = Cm(15.0)
    t.columns[1].width = Cm(1.0)
    left, right = t.rows[0].cells
    cell_text(left, ("    " * indent) + title)
    number = key if key in "กขคงจฉชซฌญฎฏ" and len(key) == 1 else PAGE_MAP.get(key, "…")
    cell_text(right, str(number), align=WD_ALIGN_PARAGRAPH.RIGHT)
    for cell in (left, right):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0


def main_chapter(number, title):
    if number > 1:
        page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"บทที่ {number}")
    p2 = doc.add_paragraph(style="Heading 1")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run(title)
    return p2


def field_page(p):
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, display, end):
        run._r.append(element)
    run.font.name = FONT
    run.font.size = Pt(16)


# ปกและส่วนต้นเล่ม
spacer(45)
paragraph("(นำภาพตรามหาวิทยาลัยเกษตรศาสตร์มาติด)", align=WD_ALIGN_PARAGRAPH.CENTER)
spacer(45)
paragraph("รายงานการปฏิบัติงานสหกิจศึกษา", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
spacer(28)
paragraph("เรื่อง", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
spacer(25)
paragraph("Document Management System (DMS)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
paragraph("ระบบจัดการเอกสารและอนุมัติเอกสารอิเล็กทรอนิกส์", align=WD_ALIGN_PARAGRAPH.CENTER)
spacer(45)
paragraph("นายธนภัทร ชูสุวรรณ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
paragraph("รหัสนิสิต 6521601857", align=WD_ALIGN_PARAGRAPH.CENTER)
spacer(65)
paragraph("สาขาวิชาวิทยาการคอมพิวเตอร์ คณะศิลปศาสตร์และวิทยาศาสตร์", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
paragraph("มหาวิทยาลัยเกษตรศาสตร์ วิทยาเขตกำแพงแสน ปีการศึกษา 2569", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)

page_break()
spacer(120)
paragraph("รายงานการปฏิบัติงานสหกิจศึกษา", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
spacer(30)
paragraph("เรื่อง", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
spacer(20)
paragraph("Document Management System (DMS)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
paragraph("ระบบจัดการเอกสารและอนุมัติเอกสารอิเล็กทรอนิกส์", align=WD_ALIGN_PARAGRAPH.CENTER)
spacer(25)
paragraph("โดย", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
spacer(35)
paragraph("นายธนภัทร ชูสุวรรณ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
paragraph("รหัสนิสิต 6521601857", align=WD_ALIGN_PARAGRAPH.CENTER)
spacer(115)
paragraph("สาขาวิชาวิทยาการคอมพิวเตอร์ คณะศิลปศาสตร์และวิทยาศาสตร์", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
paragraph("มหาวิทยาลัยเกษตรศาสตร์ วิทยาเขตกำแพงแสน ปีการศึกษา 2569", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)

top_letter("ก")
paragraph("สถานที่ 528/89 ซอยลาดพร้าว 126 แขวงพลับพลา เขตวังทองหลาง กรุงเทพมหานคร 10310", align=WD_ALIGN_PARAGRAPH.RIGHT)
paragraph("วันที่ 15 ตุลาคม 2569", align=WD_ALIGN_PARAGRAPH.RIGHT, after=9)
paragraph("เรื่อง ขอส่งรายงานการปฏิบัติงานสหกิจศึกษา", after=5)
paragraph("เรียน อาจารย์ที่ปรึกษาสหกิจศึกษา สาขาวิชาวิทยาการคอมพิวเตอร์", after=6)
body("ข้าพเจ้า นายธนภัทร ชูสุวรรณ รหัสนิสิต 6521601857 สาขาวิชาวิทยาการคอมพิวเตอร์ คณะศิลปศาสตร์และวิทยาศาสตร์ มหาวิทยาลัยเกษตรศาสตร์ วิทยาเขตกำแพงแสน ได้ปฏิบัติงานสหกิจศึกษา ณ บริษัท webpark ในตำแหน่ง Dev ระหว่างวันที่ 22 มิถุนายน 2569 ถึงวันที่ 16 ตุลาคม 2569 โดยได้รับมอบหมายงานเกี่ยวกับโครงการ Document Management System (DMS) สำหรับจัดการเอกสารและกระบวนการอนุมัติภายในองค์กร")
body("บัดนี้ขอส่งรายงานการปฏิบัติงานสหกิจศึกษาเพื่อให้อาจารย์ตรวจพิจารณา จำนวน 1 เล่ม พร้อมรายละเอียดการวิเคราะห์ ออกแบบ พัฒนา และตรวจสอบระบบตามหลักฐานที่มี")
paragraph("จึงเรียนมาเพื่อโปรดพิจารณา", before=12)
spacer(20)
paragraph("ขอแสดงความนับถือ", align=WD_ALIGN_PARAGRAPH.RIGHT)
spacer(25)
paragraph("(ลงชื่อ........................................................)", align=WD_ALIGN_PARAGRAPH.RIGHT)
paragraph("(นายธนภัทร ชูสุวรรณ)", align=WD_ALIGN_PARAGRAPH.RIGHT)
paragraph("นิสิตสหกิจศึกษา", align=WD_ALIGN_PARAGRAPH.RIGHT)

top_letter("ข", "บทคัดย่อ")
paragraph("ธนภัทร ชูสุวรรณ: Document Management System (DMS)")
paragraph("ปริญญาวิทยาศาสตรบัณฑิต (วิทยาการคอมพิวเตอร์) สาขาวิชาวิทยาการคอมพิวเตอร์")
paragraph("อาจารย์ที่ปรึกษาโครงการ: อาจารย์พิชยภรณ์ พงศกรรังศิลป์", after=8)
body("รายงานนี้นำเสนอการพัฒนาระบบจัดการเอกสารและอนุมัติเอกสารอิเล็กทรอนิกส์สำหรับบริษัท นิสซุย (ประเทศไทย) จำกัด ภายใต้การปฏิบัติงานสหกิจศึกษาที่บริษัท webpark ระบบมีเป้าหมายให้สร้างและจัดเก็บเอกสาร PR, PO, บันทึกข้อความ และ PDF กำหนดผู้อนุมัติเป็นลำดับ ติดตามสถานะ จัดเก็บเวอร์ชัน ค้นหาเอกสาร และตรวจสอบการทำงานย้อนหลังได้")
body("ส่วนติดต่อผู้ใช้พัฒนาด้วย Next.js และ TypeScript ส่วนบริการใช้ NestJS เชื่อมต่อ PostgreSQL ผ่าน Prisma ระบบประกอบด้วยการสร้างเอกสารร่าง การออกเลขเอกสาร การจัดเก็บเวอร์ชัน การส่งอนุมัติตามลำดับ การลงลายเซ็นบน PDF การแจ้งเตือน และการค้นหาเอกสารตามสิทธิ์ของผู้ใช้")
body("ผลการปฏิบัติงานได้ระบบ DMS พร้อมโครงสร้างข้อมูล หน้าจอ และบริการ API ตามขอบเขตโครงการ ผู้จัดทำตรวจสอบการทำงานของกระบวนการหลักและจัดทำเอกสารประกอบการใช้งานและการพัฒนาระบบ")
spacer(28)
paragraph("________________________     ________________________     16/10/2569", align=WD_ALIGN_PARAGRAPH.CENTER)
paragraph("ลายมือนิสิต                         ลายมืออาจารย์ที่ปรึกษา                 วันที่", align=WD_ALIGN_PARAGRAPH.CENTER)

top_letter("ค", "กิตติกรรมประกาศ")
body("การจัดทำโครงการ Document Management System (DMS) และรายงานการปฏิบัติงานสหกิจศึกษาฉบับนี้ได้รับโอกาสในการเรียนรู้กระบวนการพัฒนาระบบงานจริงจากบริษัท webpark ผู้จัดทำขอขอบคุณผู้บริหารและพนักงานที่ให้คำแนะนำเกี่ยวกับการวิเคราะห์ความต้องการ การออกแบบ การพัฒนา และการตรวจสอบระบบ")
body("ขอขอบคุณอาจารย์พิชยภรณ์ พงศกรรังศิลป์ อาจารย์ที่ปรึกษาโครงการ ที่ให้คำปรึกษาและตรวจทานรายงาน ตลอดจนผู้เกี่ยวข้องทุกท่านที่สนับสนุนการปฏิบัติงานสหกิจศึกษา")
paragraph("นายธนภัทร ชูสุวรรณ", align=WD_ALIGN_PARAGRAPH.RIGHT, before=20)
paragraph("16 ตุลาคม 2569", align=WD_ALIGN_PARAGRAPH.RIGHT)

top_letter("ง", "สารบัญ")
paragraph("หน้า", align=WD_ALIGN_PARAGRAPH.RIGHT)
for title, pg in [
    ("จดหมายขอส่งรายงานการปฏิบัติงานสหกิจศึกษา", "ก"),
    ("บทคัดย่อ", "ข"), ("กิตติกรรมประกาศ", "ค"),
    ("สารบัญ", "ง"), ("สารบัญตาราง", "ฉ"), ("สารบัญภาพ", "ช"),
    ("สารบัญภาพภาคผนวก", "ซ"),
]:
    toc_line(title, pg)
toc_line("บทที่ 1 บทนำ", "chapter1")
for label, key in [
    ("1. ที่มาและความสำคัญ", "1.1"), ("2. วัตถุประสงค์", "1.2"),
    ("3. รายละเอียดเกี่ยวกับองค์กรหรือหน่วยงาน", "1.3"),
    ("4. ขอบเขตของระบบ", "1.4"), ("5. ตำแหน่งและลักษณะงานที่ได้รับมอบหมาย", "1.5"),
    ("6. อุปกรณ์และเครื่องมือที่ใช้ในการพัฒนา", "1.6"),
]:
    toc_line(label, key, 1)
toc_line("บทที่ 2 ตรวจสอบเอกสาร", "chapter2")
for label, key in [
    ("1. ระบบจัดการเอกสารและเวอร์ชัน", "2.1"),
    ("2. ระบบอนุมัติและลายเซ็นอิเล็กทรอนิกส์", "2.2"),
    ("3. เทคโนโลยีส่วนติดต่อผู้ใช้", "2.3"),
    ("4. เทคโนโลยีบริการและฐานข้อมูล", "2.4"),
    ("5. การควบคุมสิทธิ์และประวัติการใช้งาน", "2.5"),
]:
    toc_line(label, key, 1)

top_letter("จ", "สารบัญ (ต่อ)")
paragraph("หน้า", align=WD_ALIGN_PARAGRAPH.RIGHT)
toc_line("บทที่ 3 รายละเอียดของงานที่ปฏิบัติ", "chapter3")
for label, key in [
    ("1. ศึกษาความต้องการและระบบงานเดิม", "3.1"),
    ("2. วิเคราะห์ขอบเขตและผู้ใช้งาน", "3.2"),
    ("3. ออกแบบสถาปัตยกรรมและกระบวนการ", "3.3"),
    ("4. ออกแบบกรณีใช้งาน", "3.4"),
    ("5. ออกแบบฐานข้อมูล", "3.5"),
    ("6. พัฒนาระบบ", "3.6"),
    ("7. การทดสอบและตรวจสอบ", "3.7"),
]:
    toc_line(label, key, 1)
toc_line("บทที่ 4 ผลการดำเนินงาน", "chapter4")
for label, key in [
    ("1. ผลงานตามโมดูล", "4.1"), ("2. ผลการตรวจสอบระบบ", "4.2"),
    ("3. การประเมินการใช้งาน", "4.3"),
]:
    toc_line(label, key, 1)
toc_line("บทที่ 5 สรุปผลการปฏิบัติงาน ประโยชน์ที่ได้รับ และข้อเสนอแนะ", "chapter5")
for label, key in [
    ("1. สรุปผลการดำเนินงาน", "5.1"), ("2. ผลสำเร็จตามวัตถุประสงค์", "5.2"),
    ("3. ประโยชน์ที่ได้รับ", "5.3"), ("4. ข้อเสนอแนะ", "5.4"),
]:
    toc_line(label, key, 1)
toc_line("เอกสารอ้างอิง", "references")
toc_line("ภาคผนวก ก ภาพประกอบระบบและการติดตั้ง", "appendixA")
toc_line("ภาคผนวก ข หลักฐานการทดสอบและแบบฟอร์มสหกิจ", "appendixB")
toc_line("ประวัติการศึกษา", "biography")

# รายการหน้าและคำบรรยายเติมจากข้อมูลที่สร้างจริงในภายหลัง จึงมีการเริ่มส่วนเนื้อหาก่อน
front_end = len(doc.paragraphs)

main = doc.add_section(WD_SECTION_START.NEW_PAGE)
main.page_width, main.page_height = Cm(21), Cm(29.7)
main.top_margin, main.bottom_margin = Cm(2.1), Cm(2.0)
main.left_margin, main.right_margin = Cm(2.55), Cm(2.15)
main.header_distance = Cm(1.0)
main.header.is_linked_to_previous = False
main.footer.is_linked_to_previous = False
header_p = main.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field_page(header_p)
pgnum = main._sectPr.xpath("./w:pgNumType")
if pgnum:
    pgnum[0].set(qn("w:start"), "1")
else:
    x = OxmlElement("w:pgNumType")
    x.set(qn("w:start"), "1")
    main._sectPr.append(x)

main_chapter(1, "บทนำ")
h("1. ที่มาและความสำคัญ")
body("การจัดการเอกสารด้วยไฟล์และขั้นตอนอนุมัติที่แยกจากกันทำให้ผู้เกี่ยวข้องค้นหาเอกสาร ติดตามสถานะ และตรวจสอบประวัติการเปลี่ยนแปลงได้ยาก เอกสารความต้องการของบริษัท นิสซุย (ประเทศไทย) จำกัด จึงกำหนดให้มีระบบกลางสำหรับสร้าง จัดเก็บ ค้นหา และอนุมัติเอกสารอิเล็กทรอนิกส์ พร้อมควบคุมสิทธิ์และเก็บประวัติการใช้งาน")
body("โครงการ Document Management System (DMS) พัฒนาเป็นเว็บแอปพลิเคชันสำหรับเอกสารคำขอซื้อ (PR) ใบสั่งซื้อ (PO) บันทึกข้อความ (BK) และ PDF ทั่วไป ผู้ใช้สามารถจัดทำร่าง ส่งอนุมัติ ติดตามผล และเรียกดูเวอร์ชันย้อนหลังได้ภายในระบบเดียว")
h("2. วัตถุประสงค์")
for s in ["จัดเก็บเอกสารและข้อมูลกำกับในระบบกลางที่ค้นหาและติดตามได้", "รองรับการสร้างเอกสาร PR, PO, BK และ PDF พร้อมเลขที่เอกสารอัตโนมัติ", "กำหนดลำดับผู้อนุมัติและแจ้งผลในแต่ละขั้น", "รักษาประวัติเวอร์ชัน ลายเซ็น และการกระทำที่ตรวจสอบย้อนหลังได้", "กำหนดสิทธิ์ตามบทบาทและขอบเขตการมองเห็นเอกสาร"]:
    bullet(s)
h("3. รายละเอียดเกี่ยวกับองค์กรหรือหน่วยงาน")
body("ผู้จัดทำปฏิบัติงานสหกิจศึกษาที่บริษัท webpark ในตำแหน่ง Dev ระหว่างวันที่ 22 มิถุนายน ถึง 16 ตุลาคม 2569 โดยทำงานเกี่ยวกับระบบ DMS ซึ่งมีบริษัท นิสซุย (ประเทศไทย) จำกัด เป็นหน่วยงานผู้ว่าจ้างตามเอกสารความต้องการโครงการ")
paragraph("ชื่อและที่ตั้งหน่วยงาน: บริษัท webpark, 528/89 ซอยลาดพร้าว 126 แขวงพลับพลา เขตวังทองหลาง กรุงเทพมหานคร 10310")
paragraph("พนักงานที่ปรึกษา: (กรอกชื่อ ตำแหน่ง และช่องทางติดต่อ)")
h("4. ขอบเขตของระบบ")
body("ขอบเขตหลักประกอบด้วยการจัดการเอกสารและโฟลเดอร์ แบบฟอร์ม PR/PO/BK การอัปโหลด PDF การอนุมัติหลายขั้น การลงลายเซ็นใน PDF เวอร์ชันเอกสาร การค้นหา การแจ้งเตือน แดชบอร์ด การจัดการผู้ใช้และข้อมูลหลัก โดยผลการพัฒนาแสดงในบทที่ 4")
table("ขอบเขตการทำงานตามบทบาท", ["บทบาท", "การทำงานหลัก"], [
    ["พนักงาน", "สร้างและติดตามเอกสารที่เกี่ยวข้องกับตน"],
    ["ผู้อนุมัติ", "พิจารณาเอกสารตามขั้นที่ได้รับมอบหมาย"],
    ["ผู้บริหาร", "ติดตามสถานะและข้อมูลสรุปตามสิทธิ์"],
    ["ผู้ดูแลระบบ", "จัดการผู้ใช้ สิทธิ์ ประเภทเอกสาร และลำดับอนุมัติ"],
], [3.4, 12.3])
h("5. ตำแหน่งและลักษณะงานที่ได้รับมอบหมาย")
paragraph("ตำแหน่ง: Dev")
body("ลักษณะงานในรายงานนี้ครอบคลุมการศึกษาเอกสารความต้องการ การออกแบบโครงสร้างข้อมูลและลำดับงาน การพัฒนาส่วนติดต่อและบริการ การเชื่อมฐานข้อมูล และการตรวจสอบผลตามหลักฐานของโครงการ DMS")
paragraph("ระยะเวลาปฏิบัติงาน: 22 มิถุนายน – 16 ตุลาคม 2569")
table("แผนงานการปฏิบัติงาน", ["งาน", "ช่วงดำเนินงาน", "หลักฐาน"], [
    ["ศึกษาความต้องการและออกแบบ", "22 มิ.ย. – 10 ก.ค. 2569", "Requirement, ผังระบบ"],
    ["พัฒนาหน้าจอและ API", "13 ก.ค. – 21 ส.ค. 2569", "ซอร์สโค้ด, หน้าจอ"],
    ["เชื่อมฐานข้อมูลและทดสอบ", "24 ส.ค. – 25 ก.ย. 2569", "Schema, ผลทดสอบ"],
    ["จัดทำรายงาน", "28 ก.ย. – 16 ต.ค. 2569", "รูปเล่มฉบับนี้"],
], [5.5, 4.5, 5.7])
h("6. อุปกรณ์และเครื่องมือที่ใช้ในการพัฒนา")
body("ส่วนติดต่อผู้ใช้ใช้ Next.js 16, React 19, TypeScript และ Tailwind CSS 4 ส่วนบริการใช้ NestJS 11, Prisma 5 และ PostgreSQL การจัดการ PDF ใช้ pdf-lib และการทดสอบ API/โค้ดใช้เครื่องมือที่กำหนดในโครงการ ข้อมูลเครื่องคอมพิวเตอร์และระบบปฏิบัติการของผู้จัดทำให้เติมตามอุปกรณ์ที่ใช้จริง")
paragraph("ฮาร์ดแวร์: (กรอกรุ่นคอมพิวเตอร์ หน่วยประมวลผล และหน่วยความจำที่ใช้จริง)")

main_chapter(2, "ตรวจสอบเอกสาร")
h("1. ระบบจัดการเอกสารและเวอร์ชัน")
body("ระบบจัดการเอกสารเก็บไฟล์และข้อมูลกำกับ เช่น เลขเอกสาร ชื่อ ประเภท ผู้จัดทำ สถานะ และวันที่ เอกสารความต้องการระบุการค้นหา การจัดหมวดหมู่ การดูตัวอย่าง และการดาวน์โหลด พร้อมประวัติเวอร์ชันเพื่อย้อนดูการเปลี่ยนแปลง โครงสร้างข้อมูลของโครงการปัจจุบันใช้ตาราง Document และ DocumentVersion เชื่อมกันแบบหนึ่งต่อหลาย")
h("2. ระบบอนุมัติและลายเซ็นอิเล็กทรอนิกส์")
body("กระบวนการอนุมัติกำหนดลำดับขั้นจาก Approval Matrix ของประเภทเอกสาร เมื่อส่งร่าง ระบบสร้าง Workflow และ WorkflowStep สำหรับผู้อนุมัติแต่ละขั้น การลงลายเซ็นใช้ภาพลายเซ็นวางบน PDF ภายในระบบ ซึ่งต่างจากลายมือชื่อดิจิทัลที่อาศัยใบรับรอง CA เอกสารความต้องการไม่ได้รวมการลงนามชนิด CA")
h("3. เทคโนโลยีส่วนติดต่อผู้ใช้")
body("Next.js และ React ใช้แสดงหน้าสร้างเอกสาร ศูนย์เอกสาร รายการอนุมัติ และหน้าจัดการระบบ TypeScript ช่วยกำหนดรูปแบบข้อมูลระหว่างหน้าจอและบริการ ส่วน Tailwind CSS ใช้จัดรูปแบบและรองรับขนาดจอต่าง ๆ แพ็กเกจของโครงการยืนยันเทคโนโลยีเหล่านี้")
h("4. เทคโนโลยีบริการและฐานข้อมูล")
body("NestJS รับคำขอจากเว็บ ตรวจข้อมูลและสิทธิ์ แล้วเรียก Prisma เพื่ออ่านหรือเขียน PostgreSQL ฐานข้อมูลแยกข้อมูลผู้ใช้ บทบาท ประเภทเอกสาร เอกสาร เวอร์ชัน ลำดับอนุมัติ การแจ้งเตือน และประวัติการกระทำ ปัจจุบันไฟล์ PDF บางส่วนเก็บเป็นข้อมูลไบนารีในฐานข้อมูล จึงต้องวางแผนสำรองและกู้คืนทั้งข้อมูลและไฟล์ร่วมกัน")
h("5. การควบคุมสิทธิ์และประวัติการใช้งาน")
body("ผู้ใช้เข้าสู่ระบบด้วยบัญชีและรหัสผ่าน ระบบใช้ JWT และกำหนดบทบาท Employee, Manager, Executive และ Administrator การควบคุมสิทธิ์ครอบคลุมหน้ารายละเอียด การดูตัวอย่าง การดาวน์โหลด การอนุมัติ และการเรียกดูเวอร์ชัน พร้อมบันทึกประวัติการกระทำเพื่อการตรวจสอบย้อนหลัง")

main_chapter(3, "รายละเอียดของงานที่ปฏิบัติ")
h("1. ศึกษาความต้องการและระบบงานเดิม")
body("ผู้จัดทำศึกษาข้อกำหนดระบบ DMS ของบริษัท นิสซุย (ประเทศไทย) จำกัด ซึ่งต้องการลดการใช้กระดาษ เพิ่มความเร็วในการค้นหาและอนุมัติเอกสาร และติดตามสถานะได้จากระบบกลาง ข้อกำหนดครอบคลุมเอกสาร PR, PO, บันทึกข้อความ และ PDF รวมถึงผู้ใช้ราว 100 คนและการติดตั้งภายในองค์กร")
figure("ภาพรวมกระบวนการเดิมและกระบวนการใหม่", "ผัง As-Is และ To-Be ที่จัดทำจากข้อมูลจริง")
h("2. วิเคราะห์ขอบเขตและผู้ใช้งาน")
body("กำหนดกลุ่มผู้ใช้เป็นผู้สร้างเอกสาร ผู้อนุมัติ และผู้ดูแลระบบ โดยความสามารถในการเห็นหรือเปลี่ยนข้อมูลขึ้นกับบทบาทและความเกี่ยวข้องกับเอกสาร ขอบเขตเอกสารมีสถานะ Draft, Pending, Approved, Rejected, Cancelled และ Returned ตามแบบจำลองข้อมูลปัจจุบัน")
figure("แผนภาพกรณีใช้งานระบบ DMS", "Use Case Diagram ของผู้สร้าง ผู้อนุมัติ และผู้ดูแลระบบ")
h("3. ออกแบบสถาปัตยกรรมและกระบวนการ")
body("โครงสร้างระบบแบ่งเป็นเว็บส่วนหน้า บริการ API และฐานข้อมูล PostgreSQL ผู้ใช้ส่งข้อมูลผ่านหน้าจอ Next.js ไปยัง NestJS ซึ่งตรวจข้อมูลและเรียก Prisma เมื่อสร้างเอกสาร ระบบออกเลขเอกสาร สร้างร่างและเวอร์ชัน จากนั้นเมื่อส่งอนุมัติจึงสร้างลำดับผู้อนุมัติจาก Approval Matrix และแจ้งผู้รับผิดชอบขั้นแรก")
figure("สถาปัตยกรรมระบบ", "แผนภาพ Frontend–Backend–Database ของโปรเจกต์จริง")
figure("ผังขั้นตอนการสร้างและส่งอนุมัติเอกสาร", "Flowchart การสร้างร่าง ส่งอนุมัติ และแจ้งเตือน")
body("เมื่อผู้อนุมัติพิจารณา ระบบต้องยืนยันว่าบุคคลนั้นเป็นผู้รับผิดชอบขั้นปัจจุบัน แล้วบันทึกผลพร้อมความคิดเห็น เวลา และเวอร์ชันเอกสาร หากปฏิเสธหรือส่งกลับแก้ไข ต้องรักษาหลักฐานของรอบก่อน ส่วนการอนุมัติครบทุกขั้นจึงเปลี่ยนสถานะเอกสารเป็น Approved")
figure("ผังขั้นตอนการพิจารณาอนุมัติ", "Flowchart การอนุมัติ ปฏิเสธ ส่งกลับ และอนุมัติครบขั้น")
h("4. ออกแบบกรณีใช้งาน")
body("กรณีใช้งานหลักใช้รูปแบบเดียวกับตัวอย่าง ได้แก่ รหัส ชื่อ ผู้เกี่ยวข้อง จุดประสงค์ เงื่อนไขก่อนและหลัง ขั้นตอนหลัก และกรณีทางเลือก ตารางต่อไปสรุปกรณีที่สำคัญต่อกระบวนการเอกสาร")
table("กรณีใช้งานหลักของระบบ", ["รหัสและชื่อ", "ผู้เกี่ยวข้อง", "ขั้นตอนหลัก / ผลลัพธ์"], [
    ["UC01 เข้าสู่ระบบ", "ผู้ใช้ทุกบทบาท", "กรอกบัญชี ระบบตรวจข้อมูลและเข้าสู่หน้าตามสิทธิ์"],
    ["UC02 สร้างเอกสาร", "ผู้สร้าง", "เลือกประเภท กรอกข้อมูลหรืออัปโหลด PDF แล้วบันทึก Draft"],
    ["UC03 ส่งอนุมัติ", "ผู้สร้าง", "เลือกผู้อนุมัติตาม matrix แล้วเปลี่ยนเป็น Pending"],
    ["UC04 อนุมัติ/ปฏิเสธ", "ผู้อนุมัติ", "ตรวจเอกสาร แสดงผล ติดตามขั้นถัดไปหรือจบงาน"],
    ["UC05 ค้นหาและดูเวอร์ชัน", "ผู้มีสิทธิ์", "ค้นตามข้อมูลกำกับ เปิดเอกสารและประวัติเวอร์ชัน"],
    ["UC06 จัดการข้อมูลหลัก", "ผู้ดูแลระบบ", "ตั้งค่าประเภทเอกสาร ผู้ใช้ สิทธิ์ และ Approval Matrix"],
], [4.1, 3.1, 8.5])
paragraph("กรณีทางเลือก: ข้อมูลไม่ครบ สิทธิ์ไม่พอ ผู้อนุมัติไม่ตรงขั้น ไฟล์ผิดชนิด หรือการบันทึกล้มเหลว ระบบต้องแสดงข้อผิดพลาดและไม่รายงานว่าสำเร็จ")
figure("Sequence Diagram การสร้างเอกสาร", "Sequence Diagram ระหว่างผู้ใช้ Frontend API และฐานข้อมูล")
figure("Sequence Diagram การอนุมัติและลงลายเซ็น", "Sequence Diagram ของผู้อนุมัติ API PDF และฐานข้อมูล")
h("5. ออกแบบฐานข้อมูล")
body("Prisma Schema แบ่งข้อมูลเป็นข้อมูลหลัก ผู้ใช้และสิทธิ์ เอกสารและเวอร์ชัน กระบวนการอนุมัติ การแจ้งเตือน และ Audit Log ทุกตารางหลักใช้ UUID เป็นคีย์ เอกสารหนึ่งฉบับมีหลายเวอร์ชันและมี workflow ได้หนึ่งชุด ซึ่ง workflow มีหลายขั้น")
figure("ER Diagram ระบบ DMS", "ER Diagram จาก Prisma Schema ปัจจุบัน")
table("Data Dictionary ตาราง Document", ["Attribute", "ชนิด", "คำอธิบาย"], [
    ["id", "String UUID", "รหัสเอกสาร"], ["doc_number", "String", "เลขที่เอกสารไม่ซ้ำ"],
    ["title", "String", "ชื่อเอกสาร"], ["type_id", "String UUID", "ประเภทเอกสาร"],
    ["creator_id", "String UUID", "ผู้สร้าง"], ["status", "Enum", "สถานะเอกสาร"],
    ["visibility_type", "String", "ขอบเขตการมองเห็น"], ["is_deleted", "Boolean", "สถานะลบเชิงตรรกะ"],
], [4.2, 3.3, 8.2])
table("Data Dictionary ตาราง WorkflowStep และ DocumentVersion", ["ตาราง.ฟิลด์", "ชนิด", "คำอธิบาย"], [
    ["WorkflowStep.step_order", "Int", "ลำดับขั้นอนุมัติ"],
    ["WorkflowStep.approver_id", "String UUID?", "ผู้อนุมัติที่ได้รับมอบหมาย"],
    ["WorkflowStep.status", "Enum", "ผลของขั้น"],
    ["WorkflowStep.signature_applied", "Boolean", "บันทึกว่าฝังลายเซ็นแล้ว"],
    ["DocumentVersion.version_number", "Int", "ลำดับเวอร์ชัน"],
    ["DocumentVersion.file_data", "Bytes?", "ข้อมูลไฟล์ PDF ถ้ามี"],
    ["DocumentVersion.form_data", "Json?", "ข้อมูลแบบฟอร์ม ณ เวอร์ชันนั้น"],
], [6.7, 3.2, 5.8])
h("6. พัฒนาระบบ")
body("ส่วนหน้าแบ่งหน้าจอเป็นเข้าสู่ระบบ แดชบอร์ด ศูนย์เอกสาร สร้างเอกสาร รายการอนุมัติ โฟลเดอร์ การแจ้งเตือน ประวัติเวอร์ชัน และหน้าจัดการผู้ใช้หรือข้อมูลหลัก ส่วนหลังแบ่งบริการเป็น auth, documents, workflows, folders, notifications, dashboard, admin และ users")
body("การสร้าง PR, PO, BK และ PDF บันทึกเอกสารร่างพร้อมเวอร์ชันเริ่มต้น ส่วน Approval Matrix กำหนดจำนวนขั้นตามประเภทเอกสารและปรับแก้ได้ โดยข้อมูลตัวอย่างของโครงการกำหนด PR 3 ขั้น, PO 2 ขั้น, BK 1 ขั้น และ DOC 1 ขั้น")
h("7. การทดสอบและตรวจสอบ")
body("การตรวจสอบในช่วงวันที่ 28 กันยายน – 16 ตุลาคม 2569 ครอบคลุมการสร้างเอกสาร การบันทึกร่าง การส่งอนุมัติ การพิจารณาตามลำดับ การค้นหา และการดูเวอร์ชัน โดยตรวจความสอดคล้องระหว่างหน้าจอ บริการ API และข้อมูลในฐานข้อมูล รวมทั้งตรวจการคอมไพล์และการสร้างโปรแกรมของส่วนหน้าและส่วนหลัง")
table("รายการตรวจสอบการทำงานของระบบ", ["รายการ", "สิ่งที่ตรวจสอบ"], [
    ["การสร้างเอกสาร", "ข้อมูลแบบฟอร์ม เลขเอกสาร และเวอร์ชันเริ่มต้น"],
    ["การส่งอนุมัติ", "ลำดับผู้อนุมัติ สถานะ และการแจ้งเตือน"],
    ["การพิจารณา", "ผลอนุมัติ ความคิดเห็น และลายเซ็นบน PDF"],
    ["การจัดเก็บและค้นหา", "สิทธิ์การเข้าถึง ตัวกรอง และประวัติเวอร์ชัน"],
    ["การสร้างโปรแกรม", "การตรวจ TypeScript และ build ของส่วนหน้าและส่วนหลัง"],
], [6.3, 9.4])

main_chapter(4, "ผลการดำเนินงาน")
h("1. ผลงานตามโมดูล")
body("ระบบส่วนหน้ามีหน้าสร้างเอกสาร PR, PO, BK และ PDF หน้าศูนย์เอกสาร หน้ารายละเอียดเอกสาร ประวัติเวอร์ชัน รายการอนุมัติ แดชบอร์ด การแจ้งเตือน โฟลเดอร์ และหน้าจัดการข้อมูลหลัก ภาพประกอบในส่วนนี้ให้ใช้ภาพจากระบบปัจจุบันหลังตรวจข้อมูลส่วนบุคคลและเลขเอกสารที่แสดง")
for title, instr in [
    ("หน้าเข้าสู่ระบบ", "หน้าจอเข้าสู่ระบบ DMS"),
    ("หน้าแดชบอร์ด", "หน้า Dashboard ปัจจุบัน"),
    ("หน้าศูนย์เอกสาร", "หน้ารายการเอกสารและตัวกรอง"),
    ("หน้าสร้างแบบฟอร์ม PR", "หน้าฟอร์ม PR ที่กรอกข้อมูลตัวอย่าง"),
    ("หน้าสร้างแบบฟอร์ม PO", "หน้าฟอร์ม PO ที่กรอกข้อมูลตัวอย่าง"),
    ("หน้าสร้างบันทึกข้อความ BK", "หน้าฟอร์ม BK ที่กรอกข้อมูลตัวอย่าง"),
    ("หน้ารายละเอียดเอกสารและเวอร์ชัน", "หน้ารายละเอียดเอกสารและรายการเวอร์ชัน"),
    ("หน้ารายการอนุมัติ", "หน้ารายการอนุมัติและลำดับขั้น"),
    ("หน้าจัดการข้อมูลหลัก", "หน้า Approval Matrix และเลขที่เอกสาร"),
]:
    figure(title, instr)
h("2. ผลการตรวจสอบระบบ")
body("ผลการพัฒนาจัดแสดงการทำงานของหน้าสร้างเอกสาร ศูนย์เอกสาร รายการอนุมัติ และหน้าจัดการข้อมูลหลัก การตรวจสอบกระบวนการหลักใช้ข้อมูลตัวอย่างเพื่อเปรียบเทียบผลที่แสดงบนหน้าจอกับสถานะและเวอร์ชันที่จัดเก็บในระบบ")
table("ผลลัพธ์ตามหน้าที่หลักของระบบ", ["ด้าน", "ผลลัพธ์"], [
    ["การสร้างเอกสาร", "รองรับแบบฟอร์ม PR, PO, BK และไฟล์ PDF พร้อมข้อมูลกำกับ"],
    ["การอนุมัติ", "กำหนดลำดับผู้พิจารณาตามประเภทเอกสารและแสดงสถานะ"],
    ["เวอร์ชันและลายเซ็น", "จัดเก็บเวอร์ชันและแสดงข้อมูลการลงลายเซ็นในเอกสาร"],
    ["การค้นหา", "ค้นเอกสารด้วยข้อมูลกำกับและเปิดดูรายละเอียดตามสิทธิ์"],
    ["การจัดการข้อมูลหลัก", "จัดการผู้ใช้ ประเภทเอกสาร และ Approval Matrix"],
], [4.0, 11.7])
figure("ตัวอย่างหลักฐานการทดสอบระบบ", "ภาพผลทดสอบที่แสดงรหัสกรณี ผลที่คาดหวัง และผลจริง")
h("3. การประเมินการใช้งาน")
body("ระบบ DMS ช่วยรวมงานสร้าง จัดเก็บ อนุมัติ และติดตามเอกสารไว้ในจุดเดียว ผู้ใช้สามารถค้นหาข้อมูลตามประเภท สถานะ และผู้เกี่ยวข้อง ขณะที่ผู้ดูแลระบบกำหนดข้อมูลหลักและลำดับอนุมัติได้จากหน้าจัดการ การจัดเก็บเวอร์ชันและประวัติการกระทำช่วยให้ตรวจสอบเส้นทางของเอกสารได้ชัดเจน")

main_chapter(5, "สรุปผลการปฏิบัติงาน ประโยชน์ที่ได้รับ และข้อเสนอแนะ")
h("1. สรุปผลการดำเนินงาน")
body("การปฏิบัติงานสหกิจศึกษาระหว่างวันที่ 22 มิถุนายน – 16 ตุลาคม 2569 ได้พัฒนาระบบ DMS สำหรับจัดการเอกสารหลายประเภทและกระบวนการอนุมัติ โดยมีโครงสร้างฐานข้อมูล หน้าจอ และ API รองรับการสร้างเอกสาร การส่งอนุมัติ การติดตามสถานะ และการเรียกดูประวัติย้อนหลัง")
h("2. ผลสำเร็จตามวัตถุประสงค์")
body("ระบบรวมข้อมูลเอกสารไว้ในศูนย์กลางเดียว รองรับแบบฟอร์มและ PDF พร้อมเลขเอกสารและเวอร์ชัน กำหนดผู้อนุมัติเป็นลำดับ แสดงสถานะงาน และจัดเก็บข้อมูลสำหรับตรวจสอบย้อนหลัง ผลงานที่ส่งมอบประกอบด้วยซอร์สโค้ด โครงสร้างฐานข้อมูล ภาพหน้าจอ และคู่มืออธิบายกระบวนการหลัก")
h("3. ประโยชน์ที่ได้รับ")
paragraph("3.1 ด้านหน่วยงาน")
body("หน่วยงานได้ระบบกลางสำหรับจัดการเอกสารและลำดับอนุมัติ ช่วยให้ค้นหา ติดตามสถานะ และตรวจสอบประวัติเอกสารได้สะดวกขึ้น")
paragraph("3.2 ด้านมหาวิทยาลัย")
body("รายงานแสดงการประยุกต์ความรู้ด้านการวิเคราะห์ระบบ ฐานข้อมูล เว็บแอปพลิเคชัน และการทดสอบซอฟต์แวร์กับงานองค์กร")
paragraph("3.3 ด้านนิสิต")
body("ผู้จัดทำได้ฝึกวิเคราะห์ความต้องการ ออกแบบฐานข้อมูลและลำดับงาน พัฒนาส่วนหน้าและ API และจัดทำเอกสารสรุปผลการพัฒนาอย่างเป็นระบบ")
h("4. ข้อเสนอแนะ")
for s in ["เพิ่มการค้นหาแบบเต็มข้อความเพื่อช่วยค้นเนื้อหาในเอกสารจำนวนมาก", "พัฒนารายงานสรุประยะเวลาการอนุมัติและปริมาณเอกสารรายประเภท", "เพิ่มการแจ้งเตือนผ่านช่องทางที่เหมาะกับการใช้งานขององค์กร", "ขยายรูปแบบเอกสารและแม่แบบให้รองรับหน่วยงานอื่น", "กำหนดรอบทบทวนสิทธิ์และสำรองข้อมูลอย่างสม่ำเสมอ"]:
    bullet(s)

page_break()
paragraph("เอกสารอ้างอิง", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=10)
for s in [
    "[1] บริษัท นิสซุย (ประเทศไทย) จำกัด. (2569). เอกสารสรุปความต้องการของการพัฒนาระบบจัดการเอกสารและระบบอนุมัติเอกสารอิเล็กทรอนิกส์. 26 มิถุนายน 2569.",
    "[2] โครงการ Document Management System. (2569). ข้อกำหนดทางเทคนิคของโครงการ DMS (dms-tech-spec/SKILL.md).",
    "[3] โครงการ Document Management System. (2569). Prisma Schema, รายการแพ็กเกจและซอร์สโค้ดในโฟลเดอร์ dms.",
]:
    paragraph(s, after=5)

page_break()
paragraph("ภาคผนวก", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
spacer(100)
paragraph("ภาคผนวก ก", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
paragraph("ภาพประกอบระบบและการติดตั้ง", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
page_break()
paragraph("ภาคผนวก ก ภาพประกอบระบบและการติดตั้ง", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
body("ภาพภาคผนวกนี้ให้ใช้ภาพจากสภาพแวดล้อมที่ใช้งานจริง และใส่คำอธิบายขั้นตอนใต้ภาพตามรูปแบบตัวอย่าง")
figure("โครงสร้างโฟลเดอร์โปรเจกต์", "ภาพโครงสร้าง dms/client และ dms/server", appendix=True)
figure("การตั้งค่าฐานข้อมูล", "ภาพหน้าจอ PostgreSQL/Prisma โดยปิดข้อมูลลับ", appendix=True)
figure("ตัวอย่างการเริ่มระบบ", "ภาพหน้าจอคำสั่งเริ่ม frontend และ backend ที่ทำงานสำเร็จ", appendix=True)

page_break()
paragraph("ภาคผนวก ข", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
paragraph("หลักฐานการทดสอบและแบบฟอร์มสหกิจศึกษา", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
page_break()
paragraph("ภาคผนวก ข หลักฐานการทดสอบและแบบฟอร์มสหกิจศึกษา", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
body("ภาคผนวกนี้รวบรวมหลักฐานการตรวจสอบระบบระหว่างวันที่ 28 กันยายน – 16 ตุลาคม 2569 และแบบฟอร์มสหกิจศึกษาที่เกี่ยวข้อง เพื่อประกอบการพิจารณารายงาน")
for s in [
    "(นำภาพรายงานผลทดสอบล่าสุดพร้อมหลักฐาน UAT มาติด)",
    "(นำแบบแจ้งรายละเอียดงาน ตำแหน่งงาน และพนักงานที่ปรึกษา มาติด)",
    "(นำแบบแผนปฏิบัติงานสหกิจศึกษา มาติด)",
    "(นำแบบแจ้งโครงร่างรายงานการปฏิบัติงาน มาติด)",
    "(นำแบบแจ้งยืนยันส่งรายงานและแบบแจ้งรายละเอียดการปฏิบัติงาน มาติด)",
    "(นำบันทึกการปฏิบัติงานรายสัปดาห์ มาติด)",
]:
    paragraph(s, align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=8)

page_break()
paragraph("ประวัติการศึกษา", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=12)
paragraph("ชื่อ                         นายธนภัทร ชูสุวรรณ")
paragraph("รหัสนิสิต                   6521601857")
paragraph("วันเดือนปีเกิด            (กรอกวันเดือนปีเกิด)")
paragraph("ประวัติการศึกษา         (กรอกโรงเรียนระดับประถมศึกษาและมัธยมศึกษา)")
paragraph("ปัจจุบัน                    นิสิตสาขาวิชาวิทยาการคอมพิวเตอร์")
paragraph("                               คณะศิลปศาสตร์และวิทยาศาสตร์")
paragraph("                               มหาวิทยาลัยเกษตรศาสตร์ วิทยาเขตกำแพงแสน")

# ใส่สารบัญตาราง/ภาพก่อน section ของเนื้อหาเพื่อให้ลำดับตรงกับตัวอย่าง
# การแทรกใน document.xml ก่อน sectPr เดิมทำผ่าน paragraph/table XML ที่แยกเก็บไว้
body_element = doc.element.body
section_break_p = None
for item in body_element:
    if item.tag == qn("w:p") and item.xpath("./w:pPr/w:sectPr"):
        section_break_p = item
        break
assert section_break_p is not None
insert_index = body_element.index(section_break_p)

def front_list_page(letter, heading, items, prefix):
    holder = Document()
    p = holder.add_paragraph(letter)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.page_break_before = True
    t = holder.add_paragraph(heading)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    holder.add_paragraph("หน้า").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    listing = holder.add_table(rows=len(items), cols=2)
    listing.autofit = False
    listing.columns[0].width = Cm(15.0)
    listing.columns[1].width = Cm(1.0)
    for n, caption in enumerate(items, 1):
        row = listing.rows[n-1]
        cell_text(row.cells[0], caption)
        cell_text(row.cells[1], str(PAGE_MAP.get(f"{prefix}{n}", "…")), align=WD_ALIGN_PARAGRAPH.RIGHT)
    return list(holder.element.body)[:-1]

for nodes in [
    front_list_page("ฉ", "สารบัญตาราง", TABLES, "table"),
    front_list_page("ช", "สารบัญภาพ", FIGURES, "figure"),
    front_list_page("ซ", "สารบัญภาพภาคผนวก", APP_FIGURES, "appfigure"),
]:
    for node in nodes:
        body_element.insert(insert_index, node)
        insert_index += 1

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
print("tables", len(TABLES), "figures", len(FIGURES), "appendix figures", len(APP_FIGURES))
