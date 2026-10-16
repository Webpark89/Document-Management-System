from __future__ import annotations

import copy
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT=Path(r"D:/Document Management")
WORK=ROOT/"tmp/report/appendix_install"
ASSET=WORK/"assets"
SOURCE=Path(r"C:/Users/LeP52ExPower/Desktop/รายงานสหกิจศึกษา_DMS_ธนภัทร_ครบ_Data_Dictionary_22_ตาราง.docx")
OUTPUT=ROOT/"output/documents/รายงานสหกิจศึกษา_DMS_ธนภัทร_ปรับภาคผนวก_การติดตั้งโปรแกรม.docx"
THAI=r"C:/Windows/Fonts/LeelawUI.ttf"
MONO=r"C:/Windows/Fonts/consola.ttf"

def f(size,mono=False):
    return ImageFont.truetype(MONO if mono else THAI,size)

def center(draw,xy,text,size=34):
    draw.text(xy,text,font=f(size),fill="#111111",anchor="mm")

def make_guides():
    # Cropped, unaltered area of the official download page containing its MSI control.
    node=Image.open(ASSET/"node-download.jpg")
    node.crop((80,640,1175,770)).save(ASSET/"node-installer-option.png")
    im=Image.new("RGB",(1500,630),"white");d=ImageDraw.Draw(im)
    center(d,(750,48),"ลำดับการติดตั้ง Node.js และ npm บน Windows",43)
    labels=[("1","เปิดไฟล์ .msi\nแล้วเลือก Next"),("2","ยอมรับข้อตกลง\nและเลือกตำแหน่งติดตั้ง"),("3","ติดตั้ง Node.js, npm\nและเพิ่มลงใน PATH"),
            ("4","เลือก Install\nและรอจนเสร็จ"),("5","เลือก Finish\nแล้วเปิด Terminal ใหม่"),("6","ตรวจสอบด้วยคำสั่ง\nnode -v และ npm -v")]
    for i,(num,label) in enumerate(labels):
        col=i%3;row=i//3;x=65+col*490;y=105+row*250
        d.rectangle((x,y,x+390,y+176),outline="#333333",width=3)
        center(d,(x+195,y+31),num,33)
        for j,line in enumerate(label.split("\n")):center(d,(x+195,y+84+j*43),line,31)
        if col<2:
            yy=y+88;d.line((x+402,yy,x+467,yy),fill="black",width=3);d.polygon([(x+470,yy),(x+454,yy-10),(x+454,yy+10)],fill="black")
    im.save(ASSET/"node-steps.png",dpi=(300,300))
    im=Image.new("RGB",(1500,720),"white");d=ImageDraw.Draw(im)
    center(d,(750,48),"ค่าที่ใช้เชื่อมต่อ PostgreSQL ผ่าน pgAdmin",43)
    fields=[("Name","DMS Local"),("Host name / address","localhost"),("Port","5432 (หรือพอร์ตที่ติดตั้งไว้)"),("Maintenance database","postgres"),("Username","postgres"),("Password","รหัสผ่านที่กำหนดตอนติดตั้ง")]
    for i,(label,value) in enumerate(fields):
        y=105+i*85
        d.text((78,y+12),label,font=f(32,True),fill="black")
        d.rectangle((620,y,1420,y+64),outline="#777777",width=2)
        d.text((642,y+15),value,font=f(31),fill="#222222")
    center(d,(750,670),"Register > Server  •  General / Connection  •  Save",31)
    im.save(ASSET/"pgadmin-settings.png",dpi=(300,300))
    im=Image.new("RGB",(1700,580),"white");d=ImageDraw.Draw(im)
    center(d,(850,46),"ตัวอย่างเตรียมฐานข้อมูลใหม่และไฟล์ server/.env",42)
    lines=["-- Query Tool: connect to the postgres database", "CREATE DATABASE dms_db;", "", "# server/.env",'DATABASE_URL="postgresql://postgres:<PASSWORD>@localhost:5432/dms_db?schema=public"',
           'JWT_SECRET="<YOUR_RANDOM_SECRET>"','ENCRYPTION_KEY="<64_HEX_CHARACTERS>"','FRONTEND_URL="http://localhost:3000"']
    d.rectangle((45,100,1655,540),outline="#777777",width=2)
    for i,line in enumerate(lines):d.text((70,120+i*48),line,font=f(28,True),fill="#222222")
    im.save(ASSET/"database-env-example.png",dpi=(300,300))
    im=Image.new("RGB",(1700,810),"white");d=ImageDraw.Draw(im)
    center(d,(850,46),"คำสั่งติดตั้งแพ็กเกจและเตรียม Prisma สำหรับฐานข้อมูลใหม่",40)
    sets=[(50,"Backend  (dms/server)",["cd dms\\server","npm ci","npx prisma generate","npx prisma db push"]),
          (875,"Frontend  (dms/client)",["cd dms\\client","npm ci","", "npm run dev"])]
    for x,title,lines in sets:
        d.rectangle((x,110,x+775,630),outline="#555555",width=3)
        center(d,(x+388,150),title,32)
        for i,line in enumerate(lines):d.text((x+35,210+i*74),line,font=f(34,True),fill="#151515")
    center(d,(850,690),"Backend: npm run start:dev  |  Frontend: npm run dev",31)
    center(d,(850,752),"Next.js, NestJS, TypeScript และ Prisma ติดตั้งจาก package.json ของโปรเจกต์",31)
    im.save(ASSET/"project-packages.png",dpi=(300,300))

FIGURES=[
    ("vscode-download.png","หน้าดาวน์โหลด Visual Studio Code"),
    ("vscode-license.png","การยอมรับข้อตกลงการติดตั้ง Visual Studio Code"),
    ("vscode-tasks.png","การเลือกตัวเลือกเพิ่มเติมของ Visual Studio Code"),
    ("vscode-finish.png","การติดตั้ง Visual Studio Code เสร็จสมบูรณ์"),
    ("node-installer-option.png","ตัวเลือก Windows Installer สำหรับ Node.js"),
    ("node-steps.png","ลำดับการติดตั้งและตรวจสอบ Node.js และ npm"),
    ("postgres-download.png","หน้าดาวน์โหลดตัวติดตั้ง PostgreSQL"),
    ("postgres-components.png","การเลือกส่วนประกอบของ PostgreSQL"),
    ("postgres-password.png","การกำหนดรหัสผ่านผู้ดูแล PostgreSQL"),
    ("postgres-port.png","การกำหนดพอร์ตของ PostgreSQL"),
    ("postgres-finish.png","การติดตั้ง PostgreSQL เสร็จสมบูรณ์"),
    ("pgadmin-settings.png","ตัวอย่างค่าการเชื่อมต่อ PostgreSQL ใน pgAdmin"),
    ("database-env-example.png","ตัวอย่างการสร้างฐานข้อมูลและตั้งค่าไฟล์สภาพแวดล้อม"),
    ("project-packages.png","คำสั่งติดตั้งแพ็กเกจที่ใช้พัฒนาระบบ DMS"),
]

def cell_text(cell,value):
    p=cell.paragraphs[0]
    if p.runs:
        p.runs[0].text=value
        for r in p.runs[1:]:r.text=""
    else:p.add_run(value)

def replace_text(p,value):
    if p.runs:
        p.runs[0].text=value
        for r in p.runs[1:]:r.text=""
    else:p.add_run(value)

def main():
    make_guides()
    doc=Document(SOURCE)
    old="ภาพประกอบระบบและการติดตั้ง";new="การติดตั้งโปรแกรมที่เกี่ยวข้อง"
    for p in doc.paragraphs:
        if old in p.text:replace_text(p,p.text.replace(old,new))
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if old in c.text:cell_text(c,c.text.replace(old,new))
    start=next(p for p in doc.paragraphs if p.text==f"ภาคผนวก ก {new}")
    # Replace only the existing Appendix A content. The following page break and
    # Appendix B cover/forms from the user's Desktop copy remain intact.
    nodes=[];el=start._p
    for _ in range(5):nodes.append(el);el=el.getnext()
    assert "ภาพภาคผนวกที่ ก3" in Paragraph(nodes[-1],doc).text
    marker=el
    for el in nodes:el.getparent().remove(el)
    def para(text="",bold=False,size=16,centered=False,pagebreak=False,indent=True):
        el=OxmlElement("w:p");marker.addprevious(el);p=Paragraph(el,doc)
        p.style=doc.styles['Normal']
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
        pf=p.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(7);pf.line_spacing=1.05
        pf.first_line_indent=Cm(.75) if indent and not centered else Cm(0)
        pf.page_break_before=pagebreak
        r=p.add_run(text);r.bold=bold;r.font.size=Pt(size)
        return p
    def heading(text,pagebreak=False):
        p=para(text,bold=True,pagebreak=pagebreak,indent=False);p.paragraph_format.keep_with_next=True
    def step(text):
        p=para(text);p.paragraph_format.keep_with_next=True
    def image(n,maxh=7.6,maxw=14.0):
        path,title=FIGURES[n-1];file=ASSET/path
        with Image.open(file) as im:w,h=im.size
        scale=min(maxw/w,maxh/h)
        p=para(centered=True,indent=False);p.paragraph_format.space_after=Pt(4)
        p.add_run().add_picture(str(file),width=Cm(w*scale),height=Cm(h*scale))
        p.paragraph_format.keep_with_next=True
        p=para(f"ภาพภาคผนวกที่ ก{n} {title}",size=14,centered=True,indent=False)
        p.paragraph_format.space_after=Pt(12)

    heading("1. การติดตั้งโปรแกรม Visual Studio Code (ระบบปฏิบัติการ Windows)")
    step("1.1 เข้าเว็บไซต์ https://code.visualstudio.com/download เลือก Windows และดาวน์โหลด User Installer ให้ตรงกับสถาปัตยกรรมของเครื่อง")
    image(1,maxh=5.8)
    step("1.2 เปิดไฟล์ VSCodeUserSetup ที่ดาวน์โหลด อ่านข้อตกลง เลือก “I accept the agreement” แล้วคลิก “Next”")
    image(2)

    heading("1. การติดตั้งโปรแกรม Visual Studio Code (ต่อ)",True)
    step("1.3 เลือกตำแหน่งติดตั้งและตัวเลือกเพิ่มเติม โดยเปิดใช้ “Add to PATH” และตัวเลือกเปิดโฟลเดอร์ด้วย Code จากนั้นคลิก “Next”")
    image(3)
    step("1.4 คลิก “Install” รอให้การติดตั้งเสร็จ แล้วคลิก “Finish” เพื่อเปิด Visual Studio Code ใช้งาน")
    image(4)

    heading("2. การติดตั้ง Node.js และ npm (ระบบปฏิบัติการ Windows)",True)
    step("2.1 เข้าเว็บไซต์ https://nodejs.org/en/download เลือกรุ่น LTS ที่รองรับโปรเจกต์ เลือก Windows และ x64 จากนั้นคลิก “Windows Installer (.msi)”")
    image(5,maxh=3.0)
    step("2.2 เปิดไฟล์ .msi คลิก “Next” ยอมรับข้อตกลงและเลือกตำแหน่งติดตั้ง คงส่วนประกอบ Node.js runtime, npm package manager และ Add to PATH แล้วเลือก “Install” และ “Finish” ตามลำดับ")
    image(6,maxh=7.6)
    para("2.3 เปิด Terminal ใหม่ แล้วใช้คำสั่ง node --version และ npm --version เพื่อตรวจสอบว่าระบบเรียกใช้งาน Node.js และ npm ได้")

    heading("3. การติดตั้ง PostgreSQL และ pgAdmin (ระบบปฏิบัติการ Windows)",True)
    step("3.1 เข้าเว็บไซต์ https://www.postgresql.org/download/windows/ แล้วเลือกตัวติดตั้งจาก EDB ดาวน์โหลดรุ่น Windows x86-64 และเปิดไฟล์ติดตั้ง")
    image(7,maxh=6.0)
    step("3.2 เลือกตำแหน่งติดตั้ง จากนั้นเลือก PostgreSQL Server, pgAdmin 4 และ Command Line Tools เพื่อให้มีทั้งฐานข้อมูลและเครื่องมือจัดการ แล้วคลิก “Next”")
    image(8)

    heading("3. การติดตั้ง PostgreSQL และ pgAdmin (ต่อ)",True)
    step("3.3 เลือกโฟลเดอร์เก็บข้อมูล กำหนดรหัสผ่านสำหรับผู้ใช้ postgres และยืนยันรหัสผ่าน แล้วคลิก “Next”")
    image(9)
    step("3.4 ระบุพอร์ตที่ใช้เชื่อมต่อฐานข้อมูล โดยค่าปกติคือ 5432 หากเครื่องใช้พอร์ตอื่นให้จดค่าไว้สำหรับตั้งค่าโปรเจกต์ จากนั้นคลิก “Next”")
    image(10)

    heading("3. การติดตั้ง PostgreSQL และ pgAdmin (ต่อ)",True)
    step("3.5 เลือก Locale และตรวจสอบสรุปการตั้งค่า คลิก “Next” เพื่อเริ่มติดตั้ง เมื่อเสร็จให้ยกเลิกการเปิด Stack Builder หากไม่ต้องใช้ส่วนเสริม แล้วคลิก “Finish”")
    image(11)
    heading("4. การเชื่อมต่อฐานข้อมูลด้วย pgAdmin")
    step("4.1 เปิด pgAdmin เลือก “Register > Server” ตั้งชื่อการเชื่อมต่อ แล้วกรอก Host, Port และผู้ใช้ตามที่ติดตั้ง ระบุรหัสผ่านของตนเอง จากนั้นคลิก “Save”")
    image(12,maxh=5.0)

    heading("5. การติดตั้งแพ็กเกจและเตรียมโปรเจกต์ DMS",True)
    step("5.1 สร้างฐานข้อมูลใหม่สำหรับ DMS ผ่าน Query Tool ของ pgAdmin แล้วตั้งค่าไฟล์ server/.env ให้ตรงกับฐานข้อมูล ระบุ JWT_SECRET และ ENCRYPTION_KEY ของระบบ โดยค่าภายในวงเล็บเป็นช่องสำหรับแทนค่าของเครื่องติดตั้ง")
    image(13,maxh=4.9,maxw=15.2)
    step("5.2 เปิดโฟลเดอร์โปรเจกต์ใน Visual Studio Code แล้วติดตั้งแพ็กเกจแยกใน dms/server และ dms/client ตามไฟล์ package-lock.json จากนั้นสร้าง Prisma Client และสร้างตารางในฐานข้อมูลใหม่")
    image(14,maxh=6.7,maxw=15.2)
    para("5.3 Next.js, NestJS, TypeScript และ Prisma ติดตั้งผ่านแพ็กเกจของโปรเจกต์ จึงไม่ต้องติดตั้งแยกเป็นโปรแกรมบน Windows สำหรับการสร้าง PDF ให้มี Microsoft Edge หรือ Chromium และกำหนด PDF_BROWSER_PATH เมื่อต้องระบุไฟล์เบราว์เซอร์",size=14)
    para("แหล่งภาพ: ก1–ก4 และ ก7 จากเล่มตัวอย่าง; ก5 จาก nodejs.org; ก8–ก11 จากคู่มือ EDB; ก6 และ ก12–ก14 เป็นภาพอธิบายขั้นตอนที่จัดทำประกอบรายงาน",size=11,indent=False)
    para("คู่มืออ้างอิง: code.visualstudio.com/docs/setup/windows | nodejs.org/en/download | enterprisedb.com/docs/supported-open-source/postgresql/installing/windows/ | pgadmin.org/docs/pgadmin4/latest/server_dialog.html",size=10,indent=False)

    # Keep the original front-matter styling while replacing obsolete figure entries.
    listing=doc.tables[4]
    row_template=copy.deepcopy(listing.rows[0]._tr)
    for row in list(listing.rows):listing._tbl.remove(row._tr)
    for i,(_,title) in enumerate(FIGURES,1):
        listing._tbl.append(copy.deepcopy(row_template))
        row=listing.rows[-1]
        cell_text(row.cells[0],f"ภาพภาคผนวกที่ ก{i} {title}")
        cell_text(row.cells[1],"—")
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUTPUT)
    result=Document(OUTPUT)
    assert len(result.tables)==len(Document(SOURCE).tables)
    assert sum('ภาพภาคผนวกที่ ก' in p.text for p in result.paragraphs)==14
    (WORK/'figures.json').write_text(json.dumps([title for _,title in FIGURES],ensure_ascii=False),encoding='utf-8')
    print(OUTPUT)
    print('Added 14 installation figures, replaced only Appendix A and its index.')

if __name__=='__main__':main()
