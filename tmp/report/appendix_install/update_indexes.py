from pathlib import Path
import re
from pypdf import PdfReader
from docx import Document

ROOT=Path(r'D:/Document Management')
OUT=ROOT/'output/documents/รายงานสหกิจศึกษา_DMS_ธนภัทร_ปรับภาคผนวก_การติดตั้งโปรแกรม.docx'
PDF=ROOT/'tmp/report/appendix_install/edited-preview.pdf'

def replace(cell,value):
    p=cell.paragraphs[0]
    if p.runs:
        p.runs[0].text=str(value)
        for run in p.runs[1:]:run.text=''
    else:p.add_run(str(value))

reader=PdfReader(PDF)
figs={};cover={}
for i,page in enumerate(reader.pages):
    if i<10:continue
    text=page.extract_text() or ''
    lines=[line.strip() for line in text.splitlines() if line.strip()]
    printed=int(lines[0]) if lines and lines[0].isdigit() else i-9
    for m in re.finditer(r'ภาพภาคผนวกที่\s*ก(\d+)\b',text):figs.setdefault(int(m.group(1)),printed)
    if 'ภาคผนวก ก' in text and 'การติดตั้งโปรแกรมที่เกี่ยวข้อง' in text:cover.setdefault('A',printed)
    if 'ภาคผนวก ข' in text:cover.setdefault('B',printed)
    if 'ประวัติการศึกษา' in text:cover.setdefault('history',printed)
assert set(figs)==set(range(1,15)),figs
assert set(cover)=={'A','B','history'},cover
doc=Document(OUT)
for row in doc.tables[4].rows:
    number=int(re.search(r'ก(\d+)',row.cells[0].text).group(1))
    replace(row.cells[1],figs[number])
for table in doc.tables[:2]:
    for row in table.rows:
        label=row.cells[0].text.strip()
        if label.startswith('ภาคผนวก ก'):replace(row.cells[1],cover['A'])
        elif label.startswith('ภาคผนวก ข'):replace(row.cells[1],cover['B'])
        elif label=='ประวัติการศึกษา':replace(row.cells[1],cover['history'])
doc.save(OUT)
print('Figure pages',figs)
print('Contents pages',cover)
