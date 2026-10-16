from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document

ROOT = Path(r"D:/Document Management")
SOURCE = next((ROOT / "output/documents").glob("*Data_Dictionary_22_*.docx"))
PAGES = ROOT / "tmp/report/paragraph_pages.json"

def norm(value):
    return " ".join(value.replace("\r", "").replace("\x07", "").split())

def set_cell(cell, value):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = str(value)
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(str(value))

def main():
    paragraphs = json.loads(PAGES.read_text(encoding="utf-8-sig"))
    body = [(norm(x["text"]), int(x["page"])) for x in paragraphs]
    exact = {}
    tables = {}
    figures = {}
    appendix = {}
    for label, page in body:
        exact.setdefault(label, page)
        m = re.match(r"^ตารางที่\s*(\d+)\b", label)
        if m:
            tables.setdefault(int(m.group(1)), page)
        m = re.match(r"^ภาพที่\s*(\d+)\b", label)
        if m:
            figures.setdefault(int(m.group(1)), page)
        m = re.search(r"ภาพภาคผนวกที่\s*ก(\d+)", label)
        if m:
            appendix.setdefault(int(m.group(1)), page)
    assert set(tables) == set(range(1, 28)), sorted(tables)
    assert set(figures) == set(range(1, 19)), sorted(figures)
    assert set(appendix) == {1, 2, 3}, appendix
    doc = Document(SOURCE)
    for row in doc.tables[2].rows:
        m = re.match(r"^ตารางที่\s*(\d+)", row.cells[0].text)
        assert m
        set_cell(row.cells[1], tables[int(m.group(1))])
    for row in doc.tables[3].rows:
        m = re.match(r"^ภาพที่\s*(\d+)", row.cells[0].text)
        assert m
        set_cell(row.cells[1], figures[int(m.group(1))])
    for row in doc.tables[4].rows:
        m = re.search(r"ก(\d+)", row.cells[0].text)
        assert m
        set_cell(row.cells[1], appendix[int(m.group(1))])
    for table in doc.tables[:2]:
        for row in table.rows:
            label = norm(row.cells[0].text)
            if label.startswith("บทที่"):
                chapter = re.match(r"บทที่\s*(\d+)", label)
                assert chapter
                key = f"บทที่ {chapter.group(1)}"
                assert key in exact, key
                set_cell(row.cells[1], exact[key])
            elif label.startswith("ภาคผนวก ก"):
                set_cell(row.cells[1], exact["ภาคผนวก ก"])
            elif label.startswith("ภาคผนวก ข"):
                set_cell(row.cells[1], exact["ภาคผนวก ข"])
            elif re.match(r"^\d+\.\s", label) or label in {"เอกสารอ้างอิง", "ประวัติการศึกษา"}:
                assert label in exact, label
                set_cell(row.cells[1], exact[label])
    doc.save(SOURCE)
    print(SOURCE)
    print("table pages", tables)
    print("figure pages", figures)
    print("appendix pages", appendix)

if __name__ == "__main__":
    main()
