from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(r"D:/Document Management")
SOURCE = next((ROOT / "output/documents").glob("*ER_22_*.docx"))
OUTPUT = ROOT / "output/documents/รายงานสหกิจศึกษา_DMS_ธนภัทร_ครบ_Data_Dictionary_22_ตาราง.docx"
SCHEMA = ROOT / "dms/server/prisma/schema.prisma"

# Physical PostgreSQL table order, grouped as in the current Prisma schema.
ORDER = [
    "departments", "positions", "document_types", "roles", "permissions",
    "role_permissions", "users", "documents", "pr_forms", "pr_form_items",
    "po_forms", "po_form_items", "bk_forms", "document_versions",
    "approval_matrix", "workflows", "workflow_steps", "notifications",
    "running_numbers", "audit_logs", "folders", "folder_permissions",
]

TABLE_THAI = {
    "departments": "แผนก", "positions": "ตำแหน่ง", "document_types": "ประเภทเอกสาร",
    "roles": "บทบาท", "permissions": "สิทธิ์การใช้งาน", "role_permissions": "สิทธิ์ตามบทบาท",
    "users": "ผู้ใช้งาน", "documents": "เอกสาร", "pr_forms": "ใบขอซื้อ",
    "pr_form_items": "รายการใบขอซื้อ", "po_forms": "ใบสั่งซื้อ", "po_form_items": "รายการใบสั่งซื้อ",
    "bk_forms": "บันทึกข้อความ", "document_versions": "เวอร์ชันเอกสาร",
    "approval_matrix": "กฎลำดับอนุมัติ", "workflows": "กระบวนการอนุมัติ",
    "workflow_steps": "ขั้นตอนอนุมัติ", "notifications": "การแจ้งเตือน",
    "running_numbers": "เลขเอกสารต่อเนื่อง", "audit_logs": "ประวัติกิจกรรม",
    "folders": "โฟลเดอร์", "folder_permissions": "สิทธิ์โฟลเดอร์",
}

COMMON = {
    "id": "รหัสรายการ", "name": "ชื่อ", "is_active": "สถานะเปิดใช้งาน",
    "created_at": "วันเวลาที่สร้าง", "updated_at": "วันเวลาที่ปรับปรุงล่าสุด",
    "is_deleted": "สถานะลบเชิงตรรกะ", "department_id": "รหัสแผนก",
    "position_id": "รหัสตำแหน่ง", "role_id": "รหัสบทบาท",
    "permission_id": "รหัสสิทธิ์", "document_id": "รหัสเอกสาร",
    "document_type_id": "รหัสประเภทเอกสาร", "user_id": "รหัสผู้ใช้งาน",
    "folder_id": "รหัสโฟลเดอร์", "creator_id": "รหัสผู้สร้าง",
    "approver_id": "รหัสผู้อนุมัติ", "workflow_id": "รหัสกระบวนการอนุมัติ",
    "status": "สถานะ", "description": "คำอธิบาย", "remark": "หมายเหตุ",
    "remarks": "หมายเหตุ", "total_amount": "ยอดเงินรวม",
    "step_order": "ลำดับขั้นอนุมัติ", "type_id": "รหัสประเภทเอกสาร",
}

EXTRA = {
    "positions": {"level": "ระดับตำแหน่ง"},
    "document_types": {"type_name": "ชื่อประเภทเอกสาร", "prefix": "รหัสนำหน้าเลขเอกสาร"},
    "permissions": {"module": "หมวดงานที่กำหนดสิทธิ์", "action": "การกระทำที่อนุญาต"},
    "users": {
        "username": "ชื่อบัญชีผู้ใช้งาน", "email": "อีเมลผู้ใช้งาน",
        "password_hash": "ค่าแฮชรหัสผ่าน", "first_name": "ชื่อจริง", "last_name": "นามสกุล",
        "signature_encrypted": "ข้อมูลลายเซ็นที่เข้ารหัส", "reset_token": "โทเค็นรีเซ็ตรหัสผ่าน",
        "reset_token_expiry": "เวลาหมดอายุของโทเค็นรีเซ็ต",
    },
    "documents": {
        "doc_number": "เลขเอกสารที่ไม่ซ้ำ", "title": "ชื่อเรื่องเอกสาร",
        "status": "สถานะเอกสาร", "search_text": "ข้อความสำหรับค้นหา",
        "tags": "ป้ายกำกับเอกสาร", "visibility_type": "ขอบเขตการมองเห็น",
        "visibility_departments": "รหัสแผนกที่มองเห็นได้",
        "visibility_users": "รหัสผู้ใช้ที่มองเห็นได้",
        "approved_at": "วันเวลาอนุมัติเอกสาร",
    },
    "pr_forms": {
        "requested_date": "วันที่ขอซื้อ", "required_date": "วันที่ต้องการใช้",
        "requester_id": "รหัสผู้ขอซื้อ", "purpose": "วัตถุประสงค์การขอซื้อ",
    },
    "pr_form_items": {
        "pr_form_id": "รหัสใบขอซื้อ", "item_name": "ชื่อรายการสินค้า",
        "quantity": "จำนวน", "unit": "หน่วยนับ", "unit_price": "ราคาต่อหน่วย",
        "total_price": "ราคารวมรายการ",
    },
    "po_forms": {
        "vendor_name": "ชื่อผู้ขาย", "vendor_contact": "ข้อมูลติดต่อผู้ขาย",
        "delivery_date": "วันที่ส่งมอบ", "payment_terms": "เงื่อนไขการชำระเงิน",
    },
    "po_form_items": {
        "po_form_id": "รหัสใบสั่งซื้อ", "item_name": "ชื่อรายการสินค้า",
        "quantity": "จำนวน", "unit": "หน่วยนับ", "unit_price": "ราคาต่อหน่วย",
        "vat": "อัตราภาษีมูลค่าเพิ่ม", "total_price": "ราคารวมรายการ",
    },
    "bk_forms": {"subject": "เรื่องของบันทึกข้อความ", "detail": "รายละเอียดบันทึกข้อความ"},
    "document_versions": {
        "version_number": "ลำดับเวอร์ชัน", "file_data": "ข้อมูลไฟล์แนบ",
        "file_size": "ขนาดไฟล์", "file_extension": "นามสกุลไฟล์",
        "form_data": "ข้อมูลแบบฟอร์ม ณ เวอร์ชันนั้น", "uploaded_by_id": "รหัสผู้บันทึกเวอร์ชัน",
    },
    "approval_matrix": {"required_role_id": "รหัสบทบาทผู้อนุมัติ"},
    "workflows": {"total_steps": "จำนวนขั้นอนุมัติทั้งหมด", "current_step": "ขั้นอนุมัติปัจจุบัน", "status": "สถานะกระบวนการอนุมัติ"},
    "workflow_steps": {
        "status": "ผลของขั้นอนุมัติ", "action_date": "วันเวลาที่ดำเนินการ",
        "comment": "ความเห็นของผู้อนุมัติ", "signature_applied": "สถานะการลงลายเซ็น",
        "return_to_step": "ขั้นที่ส่งกลับให้แก้ไข",
    },
    "notifications": {"message": "ข้อความแจ้งเตือน", "is_read": "สถานะการอ่าน", "read_at": "วันเวลาที่อ่าน"},
    "running_numbers": {
        "prefix": "คำนำหน้าเลขเอกสาร", "year_format": "รูปแบบปีในเลขเอกสาร",
        "current_number": "เลขลำดับล่าสุด", "padding_length": "จำนวนหลักของเลขลำดับ",
        "last_reset_year": "ปีที่รีเซ็ตลำดับล่าสุด",
    },
    "audit_logs": {
        "action": "ประเภทกิจกรรม", "module": "ส่วนงานที่เกิดกิจกรรม",
        "target_id": "รหัสรายการที่ถูกดำเนินการ", "details": "รายละเอียดการเปลี่ยนแปลง",
        "ip_address": "ที่อยู่ IP ของผู้ใช้งาน",
    },
    "folders": {
        "color": "สีที่ใช้แสดงโฟลเดอร์", "icon": "ไอคอนโฟลเดอร์",
        "visibility": "ขอบเขตการมองเห็น", "parent_id": "รหัสโฟลเดอร์แม่",
    },
    "folder_permissions": {"can_view": "อนุญาตให้ดู", "can_edit": "อนุญาตให้แก้ไข"},
}

SCALAR_TYPES = {"String", "Int", "Boolean", "DateTime", "Decimal", "Json", "Bytes",
                "DocumentStatus", "WorkflowStatus", "AuditAction", "FolderVisibility"}

def parse_schema():
    source = SCHEMA.read_text(encoding="utf-8")
    raw = re.findall(r"model\s+(\w+)\s*\{(.*?)\n\}", source, re.S)
    model_to_table = {}
    bodies = {}
    for model, body in raw:
        mm = re.search(r'@@map\("([^"]+)"\)', body)
        if mm:
            model_to_table[model] = mm.group(1)
            bodies[mm.group(1)] = body
    assert set(bodies) == set(ORDER)
    result = {}
    for table in ORDER:
        body = bodies[table]
        fk = {}
        for line in body.splitlines():
            m = re.search(r'^\s*\w+\s+(\w+)\??\s+@relation\([^)]*fields:\s*\[([^\]]+)\]', line)
            if m:
                for field in m.group(2).split(","):
                    fk[field.strip()] = model_to_table.get(m.group(1), m.group(1))
        composite = set()
        cm = re.search(r'@@id\(\[([^\]]+)\]\)', body)
        if cm:
            composite.update(x.strip() for x in cm.group(1).split(","))
        fields = []
        for line in body.splitlines():
            s = line.strip()
            if not s or s.startswith("//"):
                continue
            m = re.match(r'^(\w+)\s+([A-Za-z]+)(\[\]|\?)?(?:\s|$)', s)
            if not m or m.group(2) not in SCALAR_TYPES:
                continue
            name, base, suffix = m.group(1), m.group(2), m.group(3) or ""
            kind = ("Enum(" + base + ")") if base in {"DocumentStatus", "WorkflowStatus", "AuditAction", "FolderVisibility"} else base
            if base == "String" and ("@default(uuid())" in s or name in fk or name.endswith("_id") or name == "id"):
                kind = "String UUID"
            elif base == "Decimal":
                dm = re.search(r'@db.Decimal\((\d+),\s*(\d+)\)', s)
                kind = f"Decimal({dm.group(1)},{dm.group(2)})" if dm else "Decimal"
            if suffix:
                kind += suffix
            desc = EXTRA.get(table, {}).get(name, COMMON.get(name, f"ข้อมูล {name}"))
            if name == "id":
                desc = f"รหัส{TABLE_THAI[table]}"
            elif name == "name":
                desc = f"ชื่อ{TABLE_THAI[table]}"
            tags = []
            if "@id" in s or name in composite:
                tags.append("PK")
            if name in fk:
                tags.append(f"FK → {fk[name]}")
            if "@unique" in s:
                tags.append("ไม่ซ้ำ")
            if tags:
                desc = "; ".join(tags) + "; " + desc
            fields.append((name, kind, desc))
        result[table] = fields
    return result

def set_cell_text(cell, value):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = value
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(value)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)

def set_para_text(para, value):
    if para.runs:
        para.runs[0].text = value
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(value)

def mark_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))

def mark_no_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))

def build_table(doc, template_tbl, fields):
    el = copy.deepcopy(template_tbl)
    table = Table(el, doc)
    row_template = copy.deepcopy(table.rows[1]._tr)
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for name, kind, desc in fields:
        table._tbl.append(copy.deepcopy(row_template))
        row = table.rows[-1]
        for cell, value in zip(row.cells, (name, kind, desc)):
            set_cell_text(cell, value)
        mark_no_split(row)
    for cell, value in zip(table.rows[0].cells, ("Attribute", "ชนิด", "คำอธิบาย")):
        set_cell_text(cell, value)
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    mark_repeat_header(table.rows[0])
    return el

def main():
    data = parse_schema()
    doc = Document(SOURCE)
    body = doc.element.body
    template_tbl = doc.tables[8]._tbl
    caption = next(p for p in doc.paragraphs if p.text.startswith("ตารางที่ 4 Data Dictionary"))
    cap_el = caption._p
    first_tbl = cap_el.getnext()
    second_cap = first_tbl.getnext()
    second_tbl = second_cap.getnext()
    assert first_tbl.tag == qn("w:tbl") and second_tbl.tag == qn("w:tbl")
    assert "ตารางที่ 5 Data Dictionary" in Paragraph(second_cap, doc).text
    marker = second_tbl.getnext()
    for el in (cap_el, first_tbl, second_cap, second_tbl):
        body.remove(el)
    for n, table_name in enumerate(ORDER, start=4):
        cap = copy.deepcopy(cap_el)
        para = Paragraph(cap, doc)
        set_para_text(para, f"ตารางที่ {n} Data Dictionary ตาราง {table_name} ({TABLE_THAI[table_name]})")
        para.paragraph_format.keep_with_next = True
        marker.addprevious(cap)
        marker.addprevious(build_table(doc, template_tbl, data[table_name]))
    # Existing tables that follow the database section must continue the sequence.
    for old, new in (("ตารางที่ 6", "ตารางที่ 26"), ("ตารางที่ 7", "ตารางที่ 27")):
        for para in doc.paragraphs:
            if para.text.startswith(old) and "Data Dictionary" not in para.text:
                set_para_text(para, para.text.replace(old, new, 1))
    # Rebuild the front-matter list of tables with every new caption.
    list_table = doc.tables[2]
    assert len(list_table.rows) == 7
    row_template = copy.deepcopy(list_table.rows[3]._tr)
    while len(list_table.rows) > 0:
        list_table._tbl.remove(list_table.rows[-1]._tr)
    captions = [p.text for p in doc.paragraphs if re.match(r"^ตารางที่\s*\d+", p.text)]
    assert len(captions) == 27
    for label in captions:
        m = re.match(r"^ตารางที่\s*(\d+) Data Dictionary ตาราง ([a-z_]+)", label)
        if m:
            label = f"ตารางที่ {m.group(1)} Data Dictionary ตาราง {m.group(2)}"
        list_table._tbl.append(copy.deepcopy(row_template))
        row = list_table.rows[-1]
        set_cell_text(row.cells[0], label)
        set_cell_text(row.cells[1], "—")
        mark_no_split(row)
    assert len(list_table.rows) == 27
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    verify = Document(OUTPUT)
    assert len(verify.tables) == 32  # 12 - 2 + 22
    for table_name in ORDER:
        assert any(f"Data Dictionary ตาราง {table_name}" in p.text for p in verify.paragraphs)
    print(OUTPUT)
    print("models", len(data), "attributes", sum(len(x) for x in data.values()), "tables", len(verify.tables))

if __name__ == "__main__":
    main()
